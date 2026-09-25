"""Importe la liste de verbes du locuteur (français → soninké) dans l'inventaire.

La colonne soninké donne le nom d'action (radical + suffixe : daga-ye, xara-ŋe).
On en tire le radical ; la forme en -ni et la transitivité restent à confirmer.

    uv run python scripts/import_verbes_liste.py data/sources/verbes_francais_soninke.xlsx
"""

import sys
from pathlib import Path

import yaml
from openpyxl import load_workbook

from snk_engine.orthography import modernize

ROOT = Path(__file__).resolve().parents[1]
VERBES = ROOT / "data" / "references" / "verbes.yaml"


def main() -> None:
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "data" / "sources" / "verbes_francais_soninke.xlsx"
    ws = load_workbook(path, data_only=True)["Verbes"]
    rows = [r for r in ws.iter_rows(values_only=True)][1:]

    data = yaml.safe_load(VERBES.read_text(encoding="utf-8")) or {}
    known = {str(v.get("radical", "")).lower() for v in data.get("locuteur", [])}
    known |= {str(v.get("radical", "")).lower() for v in data.get("liste", [])}
    engine_radicals = set()
    grammar = yaml.safe_load((ROOT / "data" / "grammaire" / "soninke.yaml").read_text(encoding="utf-8"))
    for verb in grammar["verbes"].values():
        engine_radicals.add(str(verb.get("forme", "")).lower())

    added, skipped = [], 0
    for _, fr, groupe, _, snk in rows:
        if not snk or not fr or " ou " in str(snk):
            skipped += 1 if snk else 0
            continue
        nom_action = modernize(str(snk).strip())
        radical = nom_action.rsplit("-", 1)[0] if "-" in nom_action else nom_action
        if radical.lower() in known or radical.lower() in engine_radicals:
            continue
        known.add(radical.lower())
        entry = {"radical": radical, "fr": str(fr).strip(), "nom_action": nom_action,
                 "source": "liste du locuteur (2026-09-25)", "manque": "forme en -ni ; transitivité"}
        if groupe:
            entry["groupe_fr"] = str(groupe).split(" ")[0]
        added.append(entry)

    data.setdefault("liste", [])
    data["liste"] = (data.get("liste") or []) + added
    header = VERBES.read_text(encoding="utf-8").split("locuteur:")[0]
    body = yaml.safe_dump({k: data[k] for k in ("locuteur", "liste", "sources") if k in data},
                          allow_unicode=True, sort_keys=False, width=120)
    VERBES.write_text(header + body, encoding="utf-8")
    print(f"{len(added)} verbes ajoutés ({skipped} ignorés faute de forme claire)")
    print(f"inventaire : {len(data.get('locuteur', []))} + {len(data['liste'])} du locuteur, "
          f"{len(data.get('sources', []))} de la littérature")


if __name__ == "__main__":
    main()
