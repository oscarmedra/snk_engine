"""Propose une traduction pour les mots du corpus que le moteur ignore.

Le corpus est parallèle : chaque phrase soninké a sa phrase française. Un mot qui
revient toujours en face du même mot français en est très probablement la traduction.
C'est tout ce que fait ce script — compter les co-présences entre les deux langues.

Rien de ce qu'il sort n'est une vérité : ce sont des *propositions*, à valider par le
locuteur avant d'entrer dans data/grammaire/. Le score de Dice (0 à 1) dit seulement
à quel point les deux mots s'appellent l'un l'autre, jamais s'ils sont synonymes.

    uv run python scripts/aligner_lexique.py
    uv run python scripts/aligner_lexique.py --excel data/sources/lexique_a_valider.xlsx
"""

import argparse
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

from snk_engine.understand import QUESTION_WORDS, _key, load_lexicon

ROOT = Path(__file__).resolve().parents[1]
CORPUS = ROOT / "corpus" / "valide"
MOT = re.compile(r"[a-zà-ÿŋñ']+")
MIN_OCCURRENCES = 3      # en dessous, l'alignement n'est que du hasard


def mots_francais(phrase: str) -> set[str]:
    return set(MOT.findall(phrase.lower()))


def charger_paires() -> list[tuple[str, str]]:
    paires = []
    for fichier in sorted(CORPUS.glob("*.jsonl")):
        for ligne in fichier.read_text(encoding="utf-8").splitlines():
            if ligne.strip():
                entree = json.loads(ligne)
                paires.append((entree["snk"], entree["fr"]))
    return paires


def est_connu(cle: str, lex) -> bool:
    """Le moteur sait-il déjà quoi faire de ce mot ?"""
    return (cle in lex.engine.pronouns or cle in lex.verbs or cle in lex.reference_verbs
            or cle in lex.markers or cle in lex.glossary or cle in QUESTION_WORDS
            or cle in ("ke", "ku", "di"))


def aligner(paires, lex):
    presence_snk, presence_fr = Counter(), Counter()
    ensemble = defaultdict(Counter)
    exemples: dict[str, tuple[str, str]] = {}
    occurrences = Counter()

    for snk, fr in paires:
        cles = set()
        for brut in MOT.findall(snk):
            cle = _key(brut)
            # la particule se colle au mot suivant : on la détache avant de juger
            for particule in ("n'", "m'", "l'", "ŋ'", "q'"):
                if cle.startswith(particule) and len(cle) > len(particule):
                    cle = cle[len(particule):]
            if cle and not est_connu(cle, lex):
                cles.add(cle)
                occurrences[cle] += 1
                exemples.setdefault(cle, (snk, fr))
        francais = mots_francais(fr)
        for cle in cles:
            presence_snk[cle] += 1
            for mot in francais:
                ensemble[cle][mot] += 1
        for mot in francais:
            presence_fr[mot] += 1

    propositions = []
    for cle, total in occurrences.most_common():
        if total < MIN_OCCURRENCES:
            continue
        scores = {mot: 2 * n / (presence_snk[cle] + presence_fr[mot])
                  for mot, n in ensemble[cle].items() if n >= 2}
        if not scores:
            continue
        meilleurs = sorted(scores.items(), key=lambda x: -x[1])[:3]
        propositions.append({
            "snk": cle, "occurrences": total,
            "propose": meilleurs[0][0], "score": round(meilleurs[0][1], 2),
            "autres": ", ".join(f"{m} ({s:.2f})" for m, s in meilleurs[1:]),
            "exemple_snk": exemples[cle][0], "exemple_fr": exemples[cle][1],
        })
    return propositions


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--excel", type=Path, help="écrit un classeur à relire")
    args = parser.parse_args()

    lex = load_lexicon()
    paires = charger_paires()
    propositions = aligner(paires, lex)
    print(f"{len(paires)} paires lues · {len(propositions)} mots inconnus alignables "
          f"(vus au moins {MIN_OCCURRENCES} fois)\n")
    for p in propositions[:25]:
        print(f"   {p['snk']:14} {p['occurrences']:4}×  →  {p['propose']:16} {p['score']:.2f}")

    if not args.excel:
        print("\n(--excel <fichier> pour produire le classeur de validation)")
        return

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    ws = wb.active
    ws.title = "lexique à valider"
    entetes = ["mot soninké", "vu", "traduction proposée", "score", "autres pistes",
               "CORRECTION (laisser vide si c'est juste)", "exemple (soninké)", "exemple (français)"]
    ws.append(entetes)
    for cellule in ws[1]:
        cellule.font = Font(bold=True)
    for p in propositions:
        ws.append([p["snk"], p["occurrences"], p["propose"], p["score"], p["autres"],
                   "", p["exemple_snk"], p["exemple_fr"]])
    for colonne, largeur in zip("ABCDEFGH", (16, 6, 22, 7, 30, 38, 46, 46)):
        ws.column_dimensions[colonne].width = largeur
    for ligne in ws.iter_rows(min_row=2):
        for cellule in ligne:
            cellule.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    args.excel.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.excel)
    print(f"\nécrit : {args.excel}")


if __name__ == "__main__":
    main()
