"""Verse dans le corpus validé les phrases relues par le locuteur.

    uv run python scripts/import_validation.py data/questionnaires/validation_01.xlsx
    uv run python scripts/import_validation.py … --tout   # tout le fichier est validé

Une ligne corrigée entre dans le corpus sous sa forme corrigée, et la phrase
produite par le moteur est conservée à côté, pour qu'on voie ce qui a été repris.
"""

import json
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]


def main() -> None:
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    everything = "--tout" in sys.argv
    path = Path(args[0]) if args else ROOT / "data" / "questionnaires" / "validation_01.xlsx"
    ws = load_workbook(path, data_only=True)["Phrases"]
    rows = list(ws.iter_rows(values_only=True))[1:]

    out = ROOT / "corpus" / "valide" / f"{path.stem}.jsonl"
    kept = corrected = 0
    with out.open("w", encoding="utf-8") as f:
        for ident, snk, fr, ok, correction, remark in rows:
            if not ident or not snk:
                continue
            answer = (ok or "").strip().lower() if isinstance(ok, str) else ""
            if not everything and answer != "oui" and not correction:
                continue
            entry = {"id": ident, "snk": (correction or snk).strip(), "fr": fr,
                     "source": f"{path.name} (relu par le locuteur)",
                     "niveau": "phrase", "orthographe": "projet", "origine": "moteur"}
            if correction:
                entry["snk_moteur"] = snk
                corrected += 1
            if remark:
                entry["remarque"] = remark
            f.write(json.dumps(entry, ensure_ascii=False) + "\n")
            kept += 1
    print(f"{out} : {kept} paires validées ({corrected} corrigées)")
    total = sum(1 for p in (ROOT / "corpus" / "valide").glob("*.jsonl") for _ in p.open())
    print(f"corpus validé : {total} phrases")


if __name__ == "__main__":
    main()
