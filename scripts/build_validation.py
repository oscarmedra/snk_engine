"""Prépare un échantillon de phrases produites par le moteur, à faire valider.

Les phrases sont réparties sur tous les temps et tous les verbes, pour que la
relecture couvre l'ensemble des règles plutôt qu'une seule construction.

    uv run python scripts/build_validation.py [nombre]
"""

import random
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from snk_engine.conjugation import default_engine
from snk_engine.sentences import generate

ROOT = Path(__file__).resolve().parents[1]
HEADERS = ["ID", "Soninké (moteur)", "Français (moteur)", "Correct ?", "Correction", "Remarque"]
HELP = [
    "Pour chaque phrase : écris « oui » si elle se dit, « non » sinon.",
    "Si elle est fausse ou maladroite, écris la bonne phrase dans « Correction ».",
    "« Remarque » : tout ce que tu veux préciser (nuance, contexte, mot qui cloche).",
    "Une ligne laissée vide sera simplement ignorée.",
    "",
    "Ces phrases sont produites par les règles du moteur : ce sont elles que ta",
    "relecture valide ou corrige. Chaque « oui » devient une paire utilisable pour",
    "le corpus, chaque « non » corrige une règle.",
]


def main(count: int = 200) -> None:
    engine = default_engine()
    rows = list(generate(engine))
    buckets: dict[tuple[str, str], list] = defaultdict(list)
    for s in rows:
        buckets[(s.verb, s.tense)].append(s)

    rng = random.Random(12)
    for group in buckets.values():
        rng.shuffle(group)

    # tour par tour, une phrase de chaque couple verbe × temps
    sample, keys = [], sorted(buckets)
    while len(sample) < count and any(buckets[k] for k in keys):
        for key in keys:
            if buckets[key] and len(sample) < count:
                sample.append(buckets[key].pop())

    wb = Workbook()
    ws = wb.active
    ws.title = "Phrases"
    ws.append(HEADERS)
    fill = PatternFill("solid", fgColor="FFF2CC")
    for i, s in enumerate(sample, start=1):
        ws.append([f"v-{i:03d}", s.snk, s.fr, None, None, None])
        for col in (4, 5, 6):
            ws.cell(ws.max_row, col).fill = fill
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")
    for letter, width in zip("ABCDEF", (8, 40, 40, 10, 40, 34)):
        ws.column_dimensions[letter].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"

    dv = DataValidation(type="list", formula1='"oui,non"', allow_blank=True)
    dv.error, dv.errorTitle = "Écris oui ou non", "Réponse attendue"
    ws.add_data_validation(dv)
    dv.add(f"D2:D{len(sample) + 1}")

    help_ws = wb.create_sheet("Mode d'emploi")
    help_ws.column_dimensions["A"].width = 95
    for line in HELP:
        help_ws.append([line])

    out = ROOT / "data" / "questionnaires" / "validation_01.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    verbs = len({s.verb for s in sample})
    tenses = len({s.tense for s in sample})
    print(f"{out} : {len(sample)} phrases, {verbs} verbes, {tenses} temps")


if __name__ == "__main__":
    main(int(sys.argv[1]) if len(sys.argv) > 1 else 200)
