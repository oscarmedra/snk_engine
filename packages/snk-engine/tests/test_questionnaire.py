import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from snk_engine.conjugation import default_engine
from snk_engine.questionnaire import (
    Question,
    QuestionnaireError,
    analyse,
    load_series,
    read_answers,
    save_validated,
    write_questionnaire,
)

from snk_engine.conjugation import DATA_DIR

SERIE_02 = DATA_DIR / "questionnaires" / "serie_02.yaml"


def _filled(tmp_path: Path, questions: list[Question], replies: dict[str, tuple[str, str | None]]) -> Path:
    xlsx = write_questionnaire(questions, default_engine(), tmp_path / "q.xlsx")
    wb = load_workbook(xlsx)
    ws = wb["Questions"]
    for row in ws.iter_rows(min_row=2):
        if row[0].value in replies:
            row[7].value, row[8].value = replies[row[0].value]
    wb.save(xlsx)
    return xlsx


def test_serie_02_is_valid(tmp_path):
    series, questions = load_series(SERIE_02)
    assert series == "02" and len(questions) > 20
    write_questionnaire(questions, default_engine(), tmp_path / "s.xlsx")
    with pytest.raises(QuestionnaireError, match="écraser"):
        write_questionnaire(questions, default_engine(), tmp_path / "s.xlsx")


def test_analyse_statuses(tmp_path):
    engine = default_engine()
    qs = [
        Question("T-01", "c", "Il est parti.", "dagaer", "past", "3sg"),
        Question("T-02", "c", "Il est parti au marché.", "dagaer", "past", "3sg"),
        Question("T-03", "c", "Tu partiras.", "dagaer", "future", "2sg"),
        Question("T-04", "c", "Il a bu."),
        Question("T-05", "c", "Il est venu."),
    ]
    xlsx = _filled(tmp_path, qs, {
        "T-01": ("aké n'daga / a daga", None),
        "T-02": ("Aké n'daga saxa.", "saxa = marché"),
        "T-03": ("anké dagana", None),
        "T-04": ("a …", "boire"),
    })
    results = {r.answer.question.id: r for r in analyse(read_answers(xlsx, engine), engine)}
    assert results["T-01"].status == "conforme"
    assert results["T-02"].status == "conforme"
    assert results["T-03"].status == "à vérifier"
    assert "anke n'yi rini daga" in results["T-03"].expected
    assert results["T-04"].status == "nouveau"
    assert results["T-05"].status == "sans réponse"


def test_save_validated_splits_variants(tmp_path):
    engine = default_engine()
    qs = [Question("02-01", "c", "Il est parti.", "dagaer", "past", "3sg")]
    xlsx = _filled(tmp_path, qs, {"02-01": ("aké n'daga / a daga", "deux formes")})
    path = save_validated("02", read_answers(xlsx, engine), tmp_path / "corpus")
    rows = [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines()]
    assert [r["snk"] for r in rows] == ["ake n'daga", "a daga"]
    assert rows[0]["snk_saisie"] == "aké n'daga"
    assert rows[0]["fr"] == "Il est parti." and rows[0]["tense"] == "past"


def test_modernize_old_spelling():
    from snk_engine.orthography import modernize

    assert modernize("aké n'di maro n'yigué khakou okou") == "ake n'di maro n'yige xaku oku"
    assert modernize("yigueyer seféer mané ni anké gui ri") == "yigeyer sefeer mane ni anke gi ri"


def test_terms_workbook_round_trip(tmp_path):
    from snk_engine.questionnaire import load_terms, read_terms, write_terms

    series, terms = load_terms(DATA_DIR / "questionnaires" / "serie_04_termes.yaml")
    assert series == "04" and len(terms) > 50
    xlsx = write_terms(series, terms, tmp_path / "t.xlsx")
    rows = read_terms(xlsx)
    assert len(rows) == len(terms)
    assert rows[0]["id"] == "04-01" and rows[0]["sens"] is None
    with pytest.raises(QuestionnaireError, match="écraser"):
        write_terms(series, terms, xlsx)
