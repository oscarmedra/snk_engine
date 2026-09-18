"""Fichier Excel de travail pour organiser les verbes avant de les soumettre.

Une ligne = un verbe à un temps. Chaque colonne de pronom contient ce qui suit le
pronom (particule, marqueur, verbe). Une cellule vide = forme inconnue.
Une ligne sans temps enregistre seulement le verbe (sens, radical).

Le moteur ne lit pas ce fichier : les formes validées sont reportées dans
data/grammaire/soninke.yaml. `read_verbs` sert à vérifier et lire un classeur soumis.
"""

from __future__ import annotations

import unicodedata
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .conjugation import ConjugationEngine, ConjugationError

SHEET = "Verbes"
COL_VERB, COL_FR, COL_RADICAL, COL_TENSE, COL_NOTES = "Verbe", "Sens (fr)", "Radical", "Temps", "Remarques"

HELP = [
    "Une ligne = un verbe à un temps.",
    "Colonnes des pronoms : tout ce qui suit le pronom (particule, marqueur, verbe).",
    "   ex. dagaer, futur, anké → n'yi rini daga",
    "Cellule vide = forme inconnue : le moteur refusera de la produire.",
    "Ligne sans temps = verbe enregistré sans conjugaison (sens et radical seulement).",
    "Sens et radical peuvent n'être écrits que sur la première ligne du verbe.",
    "Remarques : libres (exceptions, n' / l', doutes...). Conservées dans les données.",
    "",
    "Une fois rempli, soumets le fichier : les formes validées seront ajoutées au moteur.",
]


class ExcelImportError(ValueError):
    pass


def _cell(value) -> str | None:
    if value is None:
        return None
    text = unicodedata.normalize("NFC", str(value)).strip()
    return text or None


def export_verbs(engine: ConjugationEngine, path: Path) -> Path:
    pronouns = list(engine.pronouns)
    headers = [COL_VERB, COL_FR, COL_RADICAL, COL_TENSE, *pronouns, COL_NOTES]

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append(headers)
    for v in engine.verbs.values():
        tenses = [t for t in engine.tenses if engine._rules(v, t)]
        if not tenses:
            ws.append([v.infinitive, v.french, v.form, None, *[None] * len(pronouns), None])
        for tense in tenses:
            forms = []
            for p in pronouns:
                try:
                    forms.append(engine.predicate(v.infinitive, p, tense))
                except ConjugationError:
                    forms.append(None)
            notes = "; ".join(r.notes for r in engine._rules(v, tense) if r.notes) or None
            ws.append([v.infinitive, v.french, v.form, engine.tenses[tense].label, *forms, notes])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    widths = {COL_VERB: 16, COL_FR: 16, COL_RADICAL: 14, COL_TENSE: 22, COL_NOTES: 40}
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = widths.get(h, 13)
    ws.freeze_panes = "B2"

    tense_col = ws.cell(1, headers.index(COL_TENSE) + 1).column_letter
    labels = ",".join(t.label for t in engine.tenses.values())
    dv = DataValidation(type="list", formula1=f'"{labels}"', allow_blank=True)
    dv.error, dv.errorTitle = "Choisir un temps de la liste", "Temps inconnu"
    ws.add_data_validation(dv)
    dv.add(f"{tense_col}2:{tense_col}2000")

    help_ws = wb.create_sheet("Mode d'emploi")
    help_ws.column_dimensions["A"].width = 100
    for line in HELP:
        help_ws.append([line])

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def read_verbs(path: Path, engine: ConjugationEngine) -> dict:
    """Lit le classeur et renvoie la section `verbes` du YAML."""
    ws = load_workbook(path, data_only=True)[SHEET]
    rows = ws.iter_rows(values_only=True)
    headers = [_cell(h) for h in next(rows)]
    for required in (COL_VERB, COL_FR, COL_TENSE):
        if required not in headers:
            raise ExcelImportError(f"Colonne manquante : '{required}'")
    idx = {h: i for i, h in enumerate(headers) if h}
    pronoun_cols = [h for h in headers if h in engine.pronouns]
    unknown = [h for h in headers if h and h not in engine.pronouns and h not in
               {COL_VERB, COL_FR, COL_RADICAL, COL_TENSE, COL_NOTES}]
    if unknown:
        raise ExcelImportError(f"Colonnes inconnues (pronoms non enregistrés ?) : {unknown}")

    tense_by_label = {}
    for code, t in engine.tenses.items():
        tense_by_label[code.lower()] = code
        tense_by_label[t.label.lower()] = code

    verbs: dict[str, dict] = {}
    for line, row in enumerate(rows, start=2):
        get = lambda col: _cell(row[idx[col]]) if col in idx and idx[col] < len(row) else None  # noqa: E731
        inf = get(COL_VERB)
        if not inf:
            if any(_cell(c) for c in row):
                raise ExcelImportError(f"Ligne {line} : colonne '{COL_VERB}' vide")
            continue
        entry = verbs.setdefault(inf, {})
        for col, key in ((COL_FR, "fr"), (COL_RADICAL, "forme")):
            value = get(col)
            if value and entry.get(key) not in (None, value):
                raise ExcelImportError(f"Ligne {line} : {col} de '{inf}' incohérent ('{entry[key]}' / '{value}')")
            if value:
                entry[key] = value

        tense_label = get(COL_TENSE)
        forms = {p: get(p) for p in pronoun_cols if get(p)}
        if not tense_label:
            if forms:
                raise ExcelImportError(f"Ligne {line} : formes renseignées sans temps pour '{inf}'")
            continue
        tense = tense_by_label.get(tense_label.lower())
        if tense is None:
            raise ExcelImportError(f"Ligne {line} : temps inconnu '{tense_label}'")
        conj = entry.setdefault("conjugaison", {})
        if tense in conj:
            raise ExcelImportError(f"Ligne {line} : '{inf}' au temps '{tense_label}' apparaît deux fois")
        rule: dict = {"formes": forms}
        if notes := get(COL_NOTES):
            rule["remarques"] = notes
        conj[tense] = rule

    for inf, entry in verbs.items():
        if "fr" not in entry:
            raise ExcelImportError(f"Sens (fr) manquant pour '{inf}'")
    return verbs

