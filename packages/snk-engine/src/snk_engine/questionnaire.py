"""Questionnaires de collecte : phrases françaises à traduire par le locuteur.

  1. Une série de questions est écrite en YAML (data/questionnaires/serie_XX.yaml).
  2. `write_questionnaire` produit l'Excel à remplir (Soninké + Description).
  3. `analyse` compare chaque réponse aux formes du moteur :
       conforme   — la réponse commence par une forme produite par le moteur ;
       à vérifier — le moteur prévoit autre chose (contradiction ou règle à revoir) ;
       nouveau    — le moteur ne sait pas encore produire cette phrase.
  4. `save_validated` enregistre les paires traduites dans data/corpus_valide/.
"""

from __future__ import annotations

import json
import re
import unicodedata
from dataclasses import asdict, dataclass, field
from pathlib import Path

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .conjugation import ConjugationEngine, ConjugationError
from .orthography import modernize

SHEET = "Questions"
HEADERS = ["ID", "Catégorie", "Français", "Verbe", "Temps", "Personne", "Indication", "Soninké", "Description"]
OPTIONAL = {"Indication"}
ALTERNATIVES = re.compile(r"\s*(?:/|\n)\s*")

HELP = [
    "Traduis chaque phrase française dans la colonne « Soninké » (en jaune).",
    "Plusieurs façons de dire ? Sépare-les par « / » ou un retour à la ligne.",
    "« Indication » contient mes questions ; réponds-y dans « Description ».",
    "Utilise « Description » pour expliquer : sens d'un mot, règle, exception, doute.",
    "Ne modifie pas les colonnes ID, Verbe, Temps et Personne : elles servent à l'analyse.",
    "Une phrase que tu ne sais pas traduire peut rester vide.",
]


class QuestionnaireError(ValueError):
    pass


@dataclass
class Question:
    id: str
    category: str
    french: str
    verb: str | None = None
    tense: str | None = None
    person: str | None = None
    hint: str | None = None


@dataclass
class Answer:
    question: Question
    soninke: str | None
    description: str | None


@dataclass
class Analysis:
    answer: Answer
    status: str  # conforme | à vérifier | nouveau | sans réponse
    expected: list[str] = field(default_factory=list)


def _clean(value) -> str | None:
    if value is None:
        return None
    text = unicodedata.normalize("NFC", str(value)).replace("’", "'").strip()
    return text or None


def load_series(path: Path) -> tuple[str, list[Question]]:
    data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    series = str(data["serie"])
    questions = [
        Question(
            id=f"{series}-{i:02d}",
            category=q["categorie"],
            french=q["fr"],
            verb=q.get("verbe"),
            tense=q.get("temps"),
            person=q.get("personne"),
            hint=q.get("indication"),
        )
        for i, q in enumerate(data.get("questions") or [], start=1)
    ]
    return series, questions


def check_series(questions: list[Question], engine: ConjugationEngine) -> None:
    persons = {p.person for forms in engine.pronouns.values() for p in forms}
    for q in questions:
        if q.verb and q.verb not in engine.verbs:
            raise QuestionnaireError(f"{q.id} : verbe inconnu du moteur '{q.verb}'")
        if q.tense and q.tense not in engine.tenses:
            raise QuestionnaireError(f"{q.id} : temps inconnu '{q.tense}'")
        if q.person and q.person not in persons:
            raise QuestionnaireError(f"{q.id} : personne inconnue '{q.person}'")


def write_questionnaire(
    questions: list[Question], engine: ConjugationEngine, path: Path, overwrite: bool = False
) -> Path:
    if path.exists() and not overwrite:
        raise QuestionnaireError(f"{path} existe déjà (réponses possibles) : refus d'écraser")
    check_series(questions, engine)
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append(HEADERS)
    to_fill = PatternFill("solid", fgColor="FFF2CC")
    for q in questions:
        tense = engine.tenses[q.tense].label if q.tense else None
        ws.append([q.id, q.category, q.french, q.verb, tense, q.person, q.hint, None, None])
        for col in (8, 9):
            ws.cell(ws.max_row, col).fill = to_fill
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")
    for letter, width in zip("ABCDEFGHI", (8, 16, 40, 12, 20, 10, 35, 40, 50)):
        ws.column_dimensions[letter].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "D2"

    help_ws = wb.create_sheet("Mode d'emploi")
    help_ws.column_dimensions["A"].width = 100
    for line in HELP:
        help_ws.append([line])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# --- Glossaire : termes relevés dans un texte, à expliquer par le locuteur -----

TERM_HEADERS = ["ID", "Catégorie", "Terme / expression", "Contexte", "Ce que je comprends",
                "Ma question", "Sens", "Explication"]
TERM_HELP = [
    "Colonnes jaunes à remplir : « Sens » (traduction courte) et « Explication » (règle, forme, remarque).",
    "« Ce que je comprends » est mon hypothèse : corrige-la si elle est fausse.",
    "Pour un verbe, donne si possible le radical, et l'infinitif si tu le connais.",
    "Une ligne que tu ne veux pas traiter peut rester vide.",
]


def load_terms(path: Path) -> tuple[str, list[dict]]:
    data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    series = str(data["serie"])
    return series, list(data.get("termes") or [])


def write_terms(series: str, terms: list[dict], path: Path, overwrite: bool = False) -> Path:
    if path.exists() and not overwrite:
        raise QuestionnaireError(f"{path} existe déjà (réponses possibles) : refus d'écraser")
    wb = Workbook()
    ws = wb.active
    ws.title = "Termes"
    ws.append(TERM_HEADERS)
    to_fill = PatternFill("solid", fgColor="FFF2CC")
    for i, t in enumerate(terms, start=1):
        ws.append([f"{series}-{i:02d}", t.get("categorie"), t.get("terme"), t.get("contexte"),
                   t.get("hypothese"), t.get("question"), None, None])
        for col in (7, 8):
            ws.cell(ws.max_row, col).fill = to_fill
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")
    for letter, width in zip("ABCDEFGH", (8, 14, 26, 38, 34, 38, 24, 46)):
        ws.column_dimensions[letter].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"

    help_ws = wb.create_sheet("Mode d'emploi")
    help_ws.column_dimensions["A"].width = 100
    for line in TERM_HELP:
        help_ws.append([line])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def read_terms(path: Path) -> list[dict]:
    """Relit le classeur rempli : une ligne par terme, avec Sens et Explication."""
    ws = load_workbook(path, data_only=True)["Termes"]
    rows = ws.iter_rows(values_only=True)
    headers = [_clean(h) for h in next(rows)]
    idx = {h: i for i, h in enumerate(headers) if h}
    out = []
    for row in rows:
        get = lambda h: _clean(row[idx[h]]) if h in idx and idx[h] < len(row) else None  # noqa: E731
        if get("ID"):
            out.append({h.lower().replace(" ", "_"): get(h) for h in TERM_HEADERS if h in idx})
    return out


def read_answers(path: Path, engine: ConjugationEngine) -> list[Answer]:
    ws = load_workbook(path, data_only=True)[SHEET]
    rows = ws.iter_rows(values_only=True)
    headers = [_clean(h) for h in next(rows)]
    missing = [h for h in HEADERS if h not in headers and h not in OPTIONAL]
    if missing:
        raise QuestionnaireError(f"Colonnes manquantes : {missing}")
    idx = {h: headers.index(h) for h in HEADERS if h in headers}
    tense_by_label = {t.label: code for code, t in engine.tenses.items()} | {c: c for c in engine.tenses}

    answers = []
    for row in rows:
        get = lambda h: _clean(row[idx[h]]) if h in idx and idx[h] < len(row) else None  # noqa: E731
        if not get("ID"):
            continue
        label = get("Temps")
        if label and label not in tense_by_label:
            raise QuestionnaireError(f"{get('ID')} : temps inconnu '{label}'")
        q = Question(
            id=get("ID"), category=get("Catégorie") or "", french=get("Français") or "",
            verb=modernize(get("Verbe")) if get("Verbe") else None, tense=tense_by_label.get(label) if label else None, person=get("Personne"),
        )
        answers.append(Answer(q, get("Soninké"), get("Description")))
    return answers


def _norm(text: str) -> str:
    """Forme de comparaison : orthographe modernisée, sans ponctuation ni trait d'union."""
    text = modernize(_clean(text) or "").replace("-", "")
    text = re.sub(r"[?!.,;]+", " ", text.lower())
    text = re.sub(r"'\s+", "'", text)
    return re.sub(r"\s+", " ", text).strip()


def expected_forms(q: Question, engine: ConjugationEngine) -> list[str]:
    if not (q.verb and q.tense):
        return []
    forms = []
    for form, pronouns in engine.pronouns.items():
        if q.person and all(p.person != q.person for p in pronouns):
            continue
        try:
            forms.append(engine.conjugate(q.verb, form, q.tense))
        except ConjugationError:
            continue
    return forms


def analyse(answers: list[Answer], engine: ConjugationEngine) -> list[Analysis]:
    results = []
    for a in answers:
        expected = expected_forms(a.question, engine)
        if not a.soninke:
            status = "sans réponse"
        elif not expected:
            status = "nouveau"
        else:
            exp = [_norm(e) for e in expected]
            alternatives = [_norm(x) for x in ALTERNATIVES.split(a.soninke) if x.strip()]
            ok = all(any(alt == e or alt.startswith(e + " ") for e in exp) for alt in alternatives)
            status = "conforme" if ok else "à vérifier"
        results.append(Analysis(a, status, expected))
    return results


def save_validated(series: str, answers: list[Answer], out_dir: Path) -> Path:
    """Enregistre les paires traduites (une ligne JSON par variante soninké)."""
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"serie_{series}.jsonl"
    with path.open("w", encoding="utf-8") as f:
        for a in answers:
            if not a.soninke:
                continue
            for variant in ALTERNATIVES.split(a.soninke):
                if not variant.strip():
                    continue
                row = {"id": a.question.id, "snk": modernize(variant.strip()), "snk_saisie": variant.strip(),
                       "fr": a.question.french,
                       "description": a.description, **{k: v for k, v in asdict(a.question).items()
                                                        if k in ("category", "verb", "tense", "person")}}
                f.write(json.dumps(row, ensure_ascii=False) + "\n")
    return path
